from capability_router import build_outline, route_capability
from generation_manager import artifact_path, render_artifact
from memory import ConversationMemory


def test_capability_router_keeps_normal_questions_on_chat_path():
    result = route_capability("等保三级的安全审计要求有哪些？")
    assert result["mode"] == "chat"


def test_writing_route_requires_delivery_context_before_outline():
    result = route_capability("帮我起草一份网络安全管理制度")
    assert result["mode"] == "writing"
    assert result["network_security"] is True
    assert "audience" in result["missing_fields"]
    assert result["clarification"]


def test_presentation_route_builds_confirmable_outline_after_fields_supplied():
    result = route_capability("制作网络安全培训PPT，12页", {
        "audience": "新入职员工",
        "purpose": "网络安全培训",
        "scenario": "培训",
        "scope": "网络安全通用场景",
    })
    assert result["mode"] == "presentation"
    assert result["missing_fields"] == []
    outline = build_outline(result["mode"], "制作网络安全培训PPT，12页", result["fields"])
    assert len(outline["pages"]) == 12


def test_generation_request_is_scoped_to_its_user_and_agent(tmp_path):
    memory = ConversationMemory(str(tmp_path / "generation.db"))
    first = memory.register_user("first@example.com", "Correct-Horse-14", "First")
    second = memory.register_user("second@example.com", "Correct-Horse-15", "Second")
    outline = build_outline("writing", "起草网络安全制度", {
        "document_type": "网络安全制度", "audience": "管理层",
        "purpose": "制度建设", "scope": "网络安全通用场景",
    })
    item = memory.create_generation_request(
        first["tenant_id"], first["id"], first["agent_id"], "writing",
        "起草网络安全制度", outline["fields"], outline,
    )
    assert memory.get_generation_request(
        item["id"], first["tenant_id"], first["id"], first["agent_id"],
    )
    assert memory.get_generation_request(
        item["id"], second["tenant_id"], second["id"], second["agent_id"],
    ) is None


def test_outline_can_be_edited_only_by_its_owner_before_rendering(tmp_path):
    memory = ConversationMemory(str(tmp_path / "generation.db"))
    first = memory.register_user("first@example.com", "Correct-Horse-14", "First")
    second = memory.register_user("second@example.com", "Correct-Horse-15", "Second")
    outline = build_outline("writing", "起草网络安全制度", {
        "document_type": "网络安全制度", "audience": "管理层",
        "purpose": "制度建设", "scope": "网络安全通用场景",
    })
    item = memory.create_generation_request(
        first["tenant_id"], first["id"], first["agent_id"], "writing",
        "起草网络安全制度", outline["fields"], outline,
    )
    updated = {**outline, "title": "修订后的网络安全制度", "sections": [
        {"heading": "适用范围", "points": ["待核验"]},
    ]}

    assert not memory.update_generation_outline(
        item["id"], second["tenant_id"], second["id"], second["agent_id"], updated,
    )
    assert memory.update_generation_outline(
        item["id"], first["tenant_id"], first["id"], first["agent_id"], updated,
    )
    assert memory.get_generation_request(
        item["id"], first["tenant_id"], first["id"], first["agent_id"],
    )["outline"]["title"] == "修订后的网络安全制度"


def test_local_word_and_ppt_artifacts_are_valid_packages(tmp_path):
    writing = build_outline("writing", "起草网络安全制度", {
        "document_type": "网络安全制度", "audience": "管理层",
        "purpose": "制度建设", "scope": "网络安全通用场景",
    })
    word_path = tmp_path / "draft.docx"
    render_artifact("writing", writing, word_path)
    from docx import Document
    assert Document(str(word_path)).paragraphs[0].text == writing["title"]

    presentation = build_outline("presentation", "制作网络安全培训PPT，5页", {
        "audience": "新员工", "purpose": "培训", "scenario": "培训",
        "page_count": "5", "scope": "网络安全通用场景",
    })
    ppt_path = tmp_path / "draft.pptx"
    render_artifact("presentation", presentation, ppt_path)
    from pptx import Presentation
    assert len(Presentation(str(ppt_path)).slides) == 5

    assert artifact_path("tenant", "user", "gen-1", "网络安全制度", "writing", 2).name.startswith("gen-1_v2_")


def test_generation_artifact_history_is_versioned_and_scoped(tmp_path):
    memory = ConversationMemory(str(tmp_path / "generation.db"))
    first = memory.register_user("first@example.com", "Correct-Horse-14", "First")
    second = memory.register_user("second@example.com", "Correct-Horse-15", "Second")
    outline = build_outline("writing", "起草网络安全制度", {
        "document_type": "网络安全制度", "audience": "管理层",
        "purpose": "制度建设", "scope": "网络安全通用场景",
    })
    item = memory.create_generation_request(
        first["tenant_id"], first["id"], first["agent_id"], "writing",
        "起草网络安全制度", outline["fields"], outline,
    )
    v1 = memory.add_generation_artifact(
        item["id"], first["tenant_id"], first["id"], first["agent_id"], "v1.docx",
    )
    v2 = memory.add_generation_artifact(
        item["id"], first["tenant_id"], first["id"], first["agent_id"], "v2.docx",
    )
    assert [artifact["version"] for artifact in memory.list_generation_artifacts(
        item["id"], first["tenant_id"], first["id"], first["agent_id"],
    )] == [2, 1]
    assert memory.get_generation_artifact(
        item["id"], v1["id"], second["tenant_id"], second["id"], second["agent_id"],
    ) is None
    assert v2["version"] == 2


def test_generation_observability_and_references_remain_scoped(tmp_path):
    memory = ConversationMemory(str(tmp_path / "generation.db"))
    first = memory.register_user("first@example.com", "Correct-Horse-14", "First")
    outline = build_outline("writing", "起草网络安全制度", {
        "document_type": "网络安全制度", "audience": "管理层",
        "purpose": "制度建设", "scope": "网络安全通用场景",
    })
    item = memory.create_generation_request(
        first["tenant_id"], first["id"], first["agent_id"], "writing",
        "起草网络安全制度", outline["fields"], outline,
    )
    trace = {"route": {"selected": "writing"}, "outcome": "outline_ready"}
    references = [{"source_name": "授权制度", "section": "第 1 章"}]
    assert memory.update_generation_observability(
        item["id"], first["tenant_id"], first["id"], first["agent_id"], trace, references,
    )
    stored = memory.get_generation_request(
        item["id"], first["tenant_id"], first["id"], first["agent_id"],
    )
    assert stored["trace"] == trace
    assert stored["references"] == references
