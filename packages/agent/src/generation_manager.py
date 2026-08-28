"""P5 generated artifact service: scoped outline records and local DOCX/PPTX rendering."""

from __future__ import annotations

import json
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


_ROOT = Path(__file__).resolve().parent.parent.parent.parent
_OUTPUT = _ROOT / "RAG_DATA" / "05_generated"


def _safe_name(value: str) -> str:
    name = re.sub(r"[^0-9A-Za-z一-鿿_-]+", "_", value).strip("_")
    return (name or "securenexus_artifact")[:80]


def render_artifact(mode: str, outline: dict[str, Any], output_path: Path,
                    version: int = 1, references: list[dict] | None = None) -> None:
    """Render a locally generated skeleton; no tenant content leaves the process."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    title = str(outline.get("title") or "网络安全交付物")
    fields = outline.get("fields") or {}
    references = references or []
    if mode == "writing":
        from docx import Document
        from docx.shared import Pt

        document = Document()
        document.add_heading(title, 0)
        document.add_paragraph(f"生成日期：{datetime.now().strftime('%Y-%m-%d')}")
        document.add_paragraph(f"版本：v{max(1, int(version))}")
        document.add_paragraph(f"适用范围：{fields.get('scope', '待确认')}")
        document.add_paragraph("说明：本初稿仅基于已确认的交付信息生成。法规、标准条款和组织事实必须结合授权资料复核。")
        for section in outline.get("sections") or []:
            document.add_heading(str(section.get("heading") or "章节"), level=1)
            for point in section.get("points") or []:
                document.add_paragraph(str(point), style="List Bullet")
        document.add_heading("引用与待核验事项", level=1)
        if references:
            for reference in references:
                document.add_paragraph(
                    f"{reference.get('source_name', '授权资料')}"
                    f" / {reference.get('section', '待定位')}", style="List Bullet",
                )
        else:
            document.add_paragraph("当前未绑定授权资料；法规、标准和组织事实均为待核验项。")
        for section in document.sections:
            section.footer.paragraphs[0].text = f"安枢 SecureNexus · 生成初稿 · v{max(1, int(version))}"
        for paragraph in document.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(10.5)
        document.save(output_path)
        return

    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    for index, page in enumerate(outline.get("pages") or []):
        layout = presentation.slide_layouts[0 if index == 0 else 1]
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = str(page.get("title") or title)
        body = slide.placeholders[1] if len(slide.placeholders) > 1 else None
        if body:
            frame = body.text_frame
            frame.clear()
            for point_index, point in enumerate(page.get("points") or []):
                paragraph = frame.paragraphs[0] if point_index == 0 else frame.add_paragraph()
                paragraph.text = str(point)
                paragraph.font.size = Pt(18)
        footer = slide.shapes.add_textbox(Inches(0.35), Inches(7.05), Inches(9.1), Inches(0.22))
        footer_frame = footer.text_frame
        footer_frame.text = (
            f"安枢 SecureNexus · v{max(1, int(version))} · "
            + ("引用待补充" if not references else "已绑定授权资料")
        )
        footer_frame.paragraphs[0].font.size = Pt(8)
    presentation.save(output_path)


def artifact_path(tenant_id: str, user_id: str, artifact_id: str, title: str, mode: str,
                  version: int = 1) -> Path:
    extension = ".docx" if mode == "writing" else ".pptx"
    return _OUTPUT / tenant_id / user_id / (
        f"{artifact_id}_v{max(1, int(version))}_{_safe_name(title)}{extension}"
    )
