#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""生成全局需求确认文档"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import datetime

doc = Document()

# ── 样式 ──
style = doc.styles['Normal']
style.font.name = 'Microsoft YaHei'
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
style.paragraph_format.line_spacing = 1.3

def h(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x24)
    return h

def p(text, bold=False, color=None):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    return para

def bullet(text):
    return doc.add_paragraph(text, style='List Bullet')

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    for i, hdr in enumerate(headers):
        t.rows[0].cells[i].text = hdr
    for row_data in rows:
        row = t.add_row()
        for i, val in enumerate(row_data):
            row.cells[i].text = str(val)
    return t

# ════════════════════════════════════════════════
# 标题页
# ════════════════════════════════════════════════
title = doc.add_heading('', level=0)
run = title.add_run('Prompt 评测 · 全局改造需求确认书')
run.font.color.rgb = RGBColor(0x60, 0xa5, 0xfa)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph()
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run(f'版本: v2.0 · {datetime.date.today().strftime("%Y-%m-%d")}\n').bold = True
meta.add_run('项目: 网络安全-移动运营商智能Agent\n')
meta.add_run('状态: 待审批 → 逐个实现')

doc.add_page_break()

# ════════════════════════════════════════════════
# 1. 概览
# ════════════════════════════════════════════════
h('一、改造概览', level=1)
p('本次改造围绕「Prompt 评测」Tab，将原有的「运行全部测试」功能全面升级为支持两种测试模式、AI 测试集生成、智能修复建议、版本管理的完整工作台。原有功能（弹性测试、导出 Word 报告、4 张总览卡片、维度贡献度图表、历史趋势图）全部保留不动。')

bullet('改造范围: 「运行全部测试」区域 → 全量测试 + 弹性测试 双模式')
bullet('保留不动: 弹性测试按钮 / 导出报告 / 4张卡片 / 图表 / System Prompt')
bullet('数据原则: 所有测试题迁移到 SQLite，两套模式共用同一数据源')
bullet('版本原则: 所有版本不可删除，还原前自动备份')

doc.add_page_break()

# ════════════════════════════════════════════════
# 2. 按钮布局
# ════════════════════════════════════════════════
h('二、工具栏按钮布局（已确认）', level=1)

table(
    ['顺序', '按钮', '说明'],
    [
        ['1', '🔬 全量测试', '默认选中，亮蓝色高亮；点击切换到全量测试模式'],
        ['2', '🔄 弹性测试', '非选中状态；点击切换到弹性测试模式'],
        ['3', '📝 提示词工作台', '弹出 SP + 版本历史面板（原名 System Prompt）'],
        ['|', '分隔线', ''],
        ['4', '📋 内置 N 条', '全量模式=31条，弹性模式=15条；文字自动切换'],
        ['5', '🎲 AI 生成', '弹窗输入关键词，生成测试题（全量20条/弹性15条）'],
        ['6', '📦 版本历史', '右侧滑出版本面板，展示所有版本+还原按钮'],
        ['7', '📄 导出 Word 报告', '生成当前测试结果的 Word 报告（挨着版本历史）'],
    ]
)

p('')
p('最终顺序: ', bold=True)
p('🔬全量测试 → 🔄弹性测试 → 📝提示词工作台 | 测试集: 📋内置 🎲AI生成 📦版本历史 📄导出Word报告')

doc.add_page_break()

# ════════════════════════════════════════════════
# 3. 全量测试
# ════════════════════════════════════════════════
h('三、🔬 全量测试模式（已确认）', level=1)

h('3.1 测试集来源', level=2)
bullet('内置 31 条: 从 prompt_test_suite.json 迁移到 SQLite，启动时自动导入')
bullet('AI 生成 20 条: 弹窗输入关键词→LLM 生成→写入 DB（set_id="ai_时间戳"）')
bullet('输入"随机"或"默认"让 AI 自主生成')

h('3.2 测试列表交互', level=2)
table(
    ['元素', '说明'],
    [
        ['#编号', '左侧灰色序号'],
        ['⏳/✅/🔴', '状态图标：待测/通过/失败'],
        ['测试问题', '截断显示，hover 可看全文'],
        ['分类标签', '来源标注/品牌禁止/越狱拦截/偏题检测/首答完整/知识准确'],
        ['难度标识', '简单/中等/困难 → 灰色/黄色/红色'],
        ['得分', '-- 或 数值'],
        ['▶ 测试按钮', '单条运行，异步评分'],
        ['✏️ 编辑按钮', '弹窗修改 query 和分类'],
        ['💡 修复按钮', '仅失败时显示→LLM生成修复建议→确认/放弃'],
    ]
)

h('3.3 全部测试（全量测试按钮）', level=2)
bullet('点击「🧪 全量测试」按钮 → 展开进度条（第X/31条 → ✅通过/🔴失败 + 分数）')
bullet('逐条异步运行，前端实时显示进度')

h('3.4 智能修复', level=2)
bullet('点击 💡 → 弹窗显示失败原因分析 + LLM 生成的 System Prompt 修改方案')
bullet('确认 → 创建新版本（v1.4）+ 覆盖 active_prompt.txt + 写 change_log')
bullet('放弃 → 丢弃建议，不做任何修改')

doc.add_page_break()

# ════════════════════════════════════════════════
# 4. 弹性测试
# ════════════════════════════════════════════════
h('四、🔄 弹性测试模式（已确认）', level=1)

p('弹性测试改造为跟全量测试完全一样的数据列表结构。', bold=True)

h('4.1 测试集来源', level=2)
bullet('内置 15 条: 覆盖口语化/复杂化/拼写错/噪音/诱导/重复等 7 种变体类型')
bullet('AI 生成 15 条: 根据关键词生成弹性变体测试题')

h('4.2 测试列表交互', level=2)
bullet('跟全量测试完全相同的 UI 布局（编号/状态/问题/分类/难度/得分/▶/✏️/💡）')
bullet('分类标签改为: 口语化/复杂化/中英混/拼写错/混入噪音/诱导/重复')
bullet('每条显示变体类型: 😐 口语 📝 加长 🎣 诱导 🌐 混合 ✖️ 错字 📢 噪音 🔁 重复')

h('4.3 评分逻辑差异', level=2)
bullet('全量测试: 6 维度评分（evaluation_matrix 权重体系）')
bullet('弹性测试: 答案一致性检查（同一问题的不同变体 → 答案内容是否一致）')

doc.add_page_break()

# ════════════════════════════════════════════════
# 5. 提示词工作台
# ════════════════════════════════════════════════
h('五、📝 提示词工作台（已确认）', level=1)

h('5.1 面板布局', level=2)
bullet('左侧: 当前 System Prompt 只读展示')
bullet('右侧: 版本历史列表 + 还原按钮')
bullet('底部提示: ⚡ 全量测试 · 弹性测试 结果均生成版本建议')

h('5.2 版本历史', level=2)
table(
    ['字段', '说明'],
    [
        ['版本号', 'v1.0, v1.1, v1.2... 自动递增'],
        ['时间', '创建时间（精确到分）'],
        ['修改人', '默认"管理员"'],
        ['修改说明', 'change_log，例如"修复数据分类缺少步骤（全量测试）"'],
        ['Diff', '显示 +N行/-N行 变更量'],
        ['是否当前', '当前版本有蓝色左边框+标记'],
        ['还原按钮', '点击→备份当前→覆盖 active_prompt.txt→Toast 提示'],
    ]
)

h('5.3 版本关联', level=2)
bullet('全量测试确认修复 → 自动创建版本，change_log 标注"（全量测试）"')
bullet('弹性测试确认修复 → 自动创建版本，change_log 标注"（弹性测试）"')
bullet('两个入口可见同一套版本数据: 提示词工作台面板 + 右侧版本面板')

doc.add_page_break()

# ════════════════════════════════════════════════
# 6. 数据层
# ════════════════════════════════════════════════
h('六、数据层设计（已确认）', level=1)

h('6.1 数据库变更', level=2)

table(
    ['表名', '操作', '说明'],
    [
        ['prompt_test_items', '新建', '题库表: set_id(builtin/ai)/seq/query/category/difficulty/is_active'],
        ['prompt_versions', '扩展', '加3字段: changed_by / change_log / prompt_diff'],
        ['version_test_results', '保留', '两套测试共用，无需改结构'],
    ]
)

h('6.2 数据迁移', level=2)
bullet('启动时检测 prompt_test_items 是否为空，为空则将 prompt_test_suite.json 的 31 条导入')
bullet('JSON 文件保留作为备份，不再作为运行时数据源')

h('6.3 System Prompt 动态加载', level=2)
bullet('新增 agent_data/active_prompt.txt — 运行时 System Prompt 源文件')
bullet('agent.py 的 SYSTEM_PROMPT 改为从 active_prompt.txt 读取')
bullet('SystemPromptLoader 实现: 文件 mtime 缓存 + 每次 ask() 检测变更')
bullet('版本还原 = 覆盖 active_prompt.txt，agent 下次调用自动生效')
bullet('还原前自动备份 → active_prompt.bak')

doc.add_page_break()

# ════════════════════════════════════════════════
# 7. 后端 API
# ════════════════════════════════════════════════
h('七、后端 API 端点（已确认）', level=1)

table(
    ['方法', '端点', '说明'],
    [
        ['POST', '/api/prompt/test/generate', 'AI 生成测试集（关键词+模式→20条/15条）'],
        ['GET', '/api/prompt/test/items', '获取当前测试集列表（内置/AI）'],
        ['PUT', '/api/prompt/test/items/{id}', '编辑单条测试题'],
        ['POST', '/api/prompt/test/run-single/{id}', '单条测试 ▶'],
        ['POST', '/api/prompt/test/run-all', '全部测试 ▶ SSE 进度流'],
        ['POST', '/api/prompt/test/suggest-fix/{id}', 'LLM 分析失败→生成修复建议'],
        ['POST', '/api/prompt/versions/create', '确认修复→创建新版本'],
        ['POST', '/api/prompt/versions/{v}/restore', '还原版本→覆盖 active_prompt.txt'],
        ['GET', '/api/prompt/versions', '获取版本历史列表'],
        ['POST', '/api/prompt/versions/elastic', '弹性测试确认修复→创建版本'],
    ]
)

doc.add_page_break()

# ════════════════════════════════════════════════
# 8. 实施计划
# ════════════════════════════════════════════════
h('八、实施计划（待批准后执行）', level=1)

table(
    ['步骤', '内容', '涉及文件', '预估'],
    [
        ['1', '数据层: DB迁移+active_prompt.txt+SystemPromptLoader', 'memory.py, agent.py', '中'],
        ['2', '核心模块: prompt_test_manager.py', '新建文件', '大'],
        ['3', '后端 API: main.py 新增10个端点', 'main.py', '中'],
        ['4', '前端: admin.html Prompt 评测 Tab 改造', 'admin.html', '大'],
        ['5', '按钮改名+顺序+导出报告集成', 'admin.html, report生成', '小'],
        ['6', '全链路测试: 生成→测试→修复→版本→还原', '手动+脚本', '中'],
    ]
)

p('')
p('注: 每个步骤完成即沟通，不批量提交。')

# ── 保存 ──
output_path = r'D:\学习资料\AI COURSE\项目\网络安全-移动运营商智能Agent\docs\全局需求确认书_v2.0.docx'
doc.save(output_path)
print(f'✅ 文档已生成: {output_path}')