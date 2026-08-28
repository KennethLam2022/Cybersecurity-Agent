from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DELIVERY = Path(r"D:\projects\项目PRD等文档\安枢SecureNexus")
DOC_NAMES = [f"{i:02d}_{name}.md" for i, name in [
    (1, "用户调研报告"), (2, "需求管理说明书"), (3, "BRD业务需求文档"),
    (4, "PRD产品需求文档"), (5, "RFC立项评审"), (6, "系统架构与详细设计文档"),
    (7, "核心流程与建模图表"), (8, "API接口文档"), (9, "软件实施计划"),
    (10, "环境配置说明"), (11, "测试策略与用例"), (12, "部署与运维手册"),
    (13, "项目README"), (14, "教师操作手册"), (15, "管理员操作手册"),
]]

COMMON = (
    "\n\n## 部署适配补充（2026-08-24）\n\n"
    "安枢 SecureNexus 已纳入 Linux 与 Docker 部署支持。生产部署统一使用独立服务账号，"
    "关闭 `ALLOW_LEGACY_LOCAL_WORKSPACE`，并通过环境变量或受控密钥系统提供模型密钥。"
    "Docker 方案提供 `Dockerfile`、`docker-compose.yml`、数据卷和容器健康检查；"
    "非容器 Linux 方案提供 `systemd` 服务文件、Nginx SSE 反向代理配置和部署自检脚本。\n\n"
    "Linux 环境不使用 Microsoft Word COM 转换旧 `.doc` 文件；文档入库应优先使用 `.docx`、PDF 或文本，"
    "旧 `.doc` 文件需在入库前转换为 `.docx`，或另行部署 LibreOffice 无头转换适配器。"
)

IDENTITY_COMMON = (
    "\n\n## 身份、登录与首次初始化补充（2026-08-25）\n\n"
    "前台与后台均采用正式浏览器会话登录：未登录访问用户端会转至 `/login`，未登录访问管理页面会保留原目标地址并转至登录页；"
    "管理接口在服务端同时校验真实登录会话、RBAC 权限、租户范围和 CSRF，不再允许本地兼容身份进入管理功能。"
    "\n\n"
    "全新实例在尚未初始化平台管理员时，访问 `/`、`/chat` 或登录页会自动进入 `/setup`。"
    "首次管理员填写显示名称、邮箱、密码和确认密码后，系统激活默认平台工作区并自动登录；"
    "初始化成功后该入口自动关闭。首次部署应先在受控内网、VPN 或防火墙限制下完成初始化，再开放外部访问。"
    "\n\n"
    "账号密码策略由服务端统一执行，覆盖首次初始化、普通注册、邀请注册和修改密码：至少 8 位，"
    "且必须同时包含大写字母、小写字母和数字；注册与首次初始化页面均要求两次密码输入一致。"
)

EXTRA = {
    "01_用户调研报告.md": "部署形态不改变用户角色、工作区隔离和审计要求；40—60 人规模建议优先采用 Linux + Docker，并由反向代理统一提供 HTTPS。",
    "02_需求管理说明书.md": "Linux/Docker 部署适配已作为已确认的交付工程项，不改变 P4—P8 功能范围；实机联调、证书配置和备份恢复仍属于上线确认。",
    "03_BRD业务需求文档.md": "企业部署应支持本地服务器、Linux 虚拟机和 Docker 环境，满足数据留存、权限隔离、可恢复和可审计要求。",
    "04_PRD产品需求文档.md": "部署配置需支持可配置监听地址、生产安全变量、健康检查、数据卷持久化和容器非 root 运行。",
    "05_RFC立项评审.md": "部署决策采用 Python/FastAPI 服务镜像加可选 Nginx；单 worker 保持 FAISS 与本地 SQLite 的稳定性，横向扩展需另行评审。",
    "06_系统架构与详细设计文档.md": "运行时目录划分为 `agent_data`、`RAG_DATA`、`upload_staging` 和 `backups`，容器和 systemd 方案均保持相同数据边界。",
    "07_核心流程与建模图表.md": "部署流程新增：构建镜像或创建虚拟环境 → 注入生产变量 → 启动服务 → 健康检查 → Nginx/入口验证 → 备份恢复演练。",
    "08_API接口文档.md": "健康检查接口为 `/api/stats/health`；容器探针和运维巡检使用该接口，不绕过认证访问管理接口。",
    "09_软件实施计划.md": "新增 Linux/Docker 实施任务：镜像构建、宿主目录权限、服务账号、反向代理、TLS、日志和恢复演练。",
    "10_环境配置说明.md": "支持 Windows、Linux 和 Docker；生产必须设置 `APP_ENV=production`、`APP_HOST`、`APP_PORT` 和 `ALLOW_LEGACY_LOCAL_WORKSPACE=0`。",
    "11_测试策略与用例.md": "新增部署验证：Linux 依赖安装、容器构建、健康检查、非 root、数据卷持久化、SSE 长连接和服务重启恢复。",
    "12_部署与运维手册.md": "已提供 Docker Compose、Linux 直接运行、systemd 和 Nginx 四种部署路径；实机验收仍需在目标服务器完成。",
    "13_项目README.md": "快速开始同时提供 Windows、Linux 和 Docker 入口；Docker 部署默认使用生产安全配置和宿主机数据卷。",
    "14_教师操作手册.md": "培训演示环境可使用 Docker 快速启动；正式培训数据应使用独立工作区和独立数据卷，避免与生产资料混用。",
    "15_管理员操作手册.md": "管理员上线检查增加容器状态、systemd 状态、健康接口、磁盘容量、数据卷权限、备份和回滚检查。",
}

IDENTITY_EXTRA = {
    "01_用户调研报告.md": "用户侧不再需要理解本地兼容工作区；登录后仅能看到已授权的工作区、Agent、资料和会话，管理员通过正式账号承担平台治理责任。",
    "02_需求管理说明书.md": "身份入口闭环和首次安装流程属于已实现的安全缺陷修复；密码规则、登录回跳、初始化关闭和未授权拒绝均应纳入回归验收。",
    "03_BRD业务需求文档.md": "业务验收要求：用户、组织管理员和平台管理员均以独立账号登录，匿名访问不得获得问答会话或后台治理权限。",
    "04_PRD产品需求文档.md": "前台登录、后台登录和首次安装共享同一身份体系；`/setup` 只在未存在已激活平台管理员时可用，创建成功后不可重复初始化。",
    "05_RFC立项评审.md": "首次安装不采用可预测默认账号或默认密码；以一次性状态判断开启初始化页，部署网络边界负责防止公网首次访问抢占管理员。",
    "06_系统架构与详细设计文档.md": "身份层由 HttpOnly 会话 Cookie、CSRF Cookie、Principal、RBAC 和租户/Agent 上下文构成；路由层负责浏览器跳转，接口层仍独立执行授权。",
    "07_核心流程与建模图表.md": "身份流程：冷启动 → `/setup` → 创建平台管理员 → 自动登录 → 后台配置；日常访问：目标页 → `/login?next=...` → 会话建立 → 按角色进入授权页面。",
    "08_API接口文档.md": "新增首次安装接口：`GET /api/auth/bootstrap-status` 返回是否需要初始化，`POST /api/auth/bootstrap` 创建首个管理员并建立会话；成功后再次调用返回冲突。",
    "09_软件实施计划.md": "安装阶段必须先限制入口网络并完成 `/setup`，再进行模型、组织、资料和权限配置；首次管理员账号由交付负责人保管并纳入交接。",
    "10_环境配置说明.md": "不再需要配置或维护首次初始化令牌；生产必须保持 `ALLOW_LEGACY_LOCAL_WORKSPACE=0`，并在首次初始化前使用网络边界限制对外访问。",
    "11_测试策略与用例.md": "新增认证用例：冷启动自动转 `/setup`、密码确认不一致拒绝、弱密码拒绝、首次初始化后入口关闭、未登录后台跳转登录、兼容身份不能访问管理接口。",
    "12_部署与运维手册.md": "Linux 与 Docker 首次启动后，在受控网络访问 `/setup` 创建平台管理员；完成后验证 `/admin` 未登录跳转、登录后权限边界和会话注销。",
    "13_项目README.md": "首次使用顺序调整为：启动服务 → 受控网络完成 `/setup` → 使用正式账号登录后台 → 配置模型与工作区；不提供默认管理员账号或默认密码。",
    "14_教师操作手册.md": "培训环境也应使用独立登录账号；讲师不应共享平台管理员密码，可由管理员预先创建培训工作区和普通成员账号。",
    "15_管理员操作手册.md": "首次初始化操作：在受控网络打开 `/setup`，填写显示名称、邮箱、两次一致的合规密码；初始化后用该平台管理员账号登录 `/admin`，再创建组织成员和角色。",
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, color in (("Title", 24, "17365D"), ("Heading 1", 16, "17365D"), ("Heading 2", 13, "1F4E79"), ("Heading 3", 11, "2F5597")):
        style = styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("安枢 SecureNexus · 文档交付版 · 2026-08-25")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(110, 110, 110)


def add_code(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F3F6F9")
    set_cell_margins(cell, 130, 160, 130, 160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("\n".join(lines))
    run.font.name = "Cascadia Mono"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    run.font.size = Pt(8.5)


def add_table(doc, rows):
    if not rows:
        return
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(rows[0]):
        cell = table.rows[0].cells[i]
        cell.text = value.strip()
        set_cell_shading(cell, "D9EAF7")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
    set_repeat_table_header(table.rows[0])
    for row in rows[1:]:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value.strip()
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    doc.add_paragraph()


def add_markdown(doc, text):
    lines = text.splitlines()
    i = 0
    in_code = False
    code = []
    table_rows = []
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            if in_code:
                add_code(doc, code)
                code = []
                in_code = False
            else:
                if table_rows:
                    add_table(doc, table_rows)
                    table_rows = []
                in_code = True
            i += 1
            continue
        if in_code:
            code.append(line)
            i += 1
            continue
        if line.strip().startswith("|") and line.strip().endswith("|"):
            parts = [p.strip() for p in line.strip().strip("|").split("|")]
            if all(set(p) <= {"-", ":", " "} for p in parts):
                i += 1
                continue
            table_rows.append(parts)
            i += 1
            continue
        if table_rows:
            add_table(doc, table_rows)
            table_rows = []
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("### "):
            doc.add_paragraph(stripped[4:], style="Heading 3")
        elif stripped.startswith("## "):
            doc.add_paragraph(stripped[3:], style="Heading 2")
        elif stripped.startswith("# "):
            doc.add_paragraph(stripped[2:], style="Title")
        elif re.match(r"^[-*] ", stripped):
            doc.add_paragraph(stripped[2:], style="List Bullet")
        elif re.match(r"^\d+\. ", stripped):
            doc.add_paragraph(re.sub(r"^\d+\. ", "", stripped), style="List Number")
        elif stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            run = p.add_run(stripped[2:])
            run.italic = True
            run.font.color.rgb = RGBColor(80, 80, 80)
        else:
            p = doc.add_paragraph()
            for n, chunk in enumerate(re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", line)):
                if not chunk:
                    continue
                run = p.add_run(chunk.strip("`") if chunk.startswith("`") else chunk.strip("*"))
                if chunk.startswith("`"):
                    run.font.name = "Cascadia Mono"
                    run.font.size = Pt(9)
                if chunk.startswith("**"):
                    run.bold = True
        i += 1
    if table_rows:
        add_table(doc, table_rows)
    if in_code and code:
        add_code(doc, code)


def update_markdown(path: Path):
    text = path.read_text(encoding="utf-8")
    marker = "## 部署适配补充（2026-08-24）"
    if marker not in text:
        text = text.rstrip() + "\n\n" + COMMON.strip() + "\n\n" + EXTRA[path.name]
    identity_marker = "## 身份、登录与首次初始化补充（2026-08-25）"
    if identity_marker not in text:
        text = text.rstrip() + "\n\n" + IDENTITY_COMMON.strip() + "\n\n" + IDENTITY_EXTRA[path.name]
    path.write_text(text.rstrip() + "\n", encoding="utf-8")
    return text


def build_docx(markdown: str, output: Path):
    doc = Document()
    configure_document(doc)
    add_markdown(doc, markdown)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main():
    if not DELIVERY.exists():
        raise SystemExit(f"delivery directory not found: {DELIVERY}")
    for name in DOC_NAMES:
        md = DELIVERY / name
        if not md.exists():
            raise SystemExit(f"missing markdown: {md}")
        markdown = update_markdown(md)
        build_docx(markdown, DELIVERY / name.replace(".md", ".docx"))
        print(f"updated {md.name}")


if __name__ == "__main__":
    main()
