"""将图片嵌入 docx — 先 add_picture 到末尾，再移动到正确位置"""
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from copy import deepcopy
from pathlib import Path
from lxml import etree

DOCX_PATH = Path(__file__).parent.parent / "docs" / "整体功能说明书_v2.0_图文完整版.docx"
IMG_DIR = Path(__file__).parent.parent / "docs" / "images"

WPML = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

INSERTIONS = [
    ("1.3 系统架构全景", "01_system_architecture.png", 5.8, "图1 网络安全移动运营商智能Agent — 系统架构图"),
    ("2.1 用户提问→回答全流程", "02_qa_flow.png", 5.5, "图2 用户提问→回答流水线（13步）"),
    ("3. RAG 检索管道（核心引擎）", "03_rag_pipeline.png", 5.8, "图3 RAG 检索流水线 — 4 阶段架构"),
    ("4. 文档预处理流水线", "04_preprocessing.png", 5.8, "图4 文档预处理流水线"),
    ("6.2 两条 LLM 调用链路", "08_llm_flow.png", 5.8, "图5 两条 LLM 调用链 — 8 张配置卡"),
    ("7. Prompt 评测体系（SRAG 质量验证）", "06_eval_framework.png", 5.8, "图6 SRAG 质量评估体系"),
    ("10.1 七层安全防护", "05_security_layers.png", 5.2, "图7 7 层安全防护体系"),
    ("11. 管理后台功能一览", "07_admin_dashboard.png", 5.8, "图8 管理后台功能总览"),
]


def find_heading_index(body, heading_text):
    """Find XML index of heading paragraph"""
    for i, child in enumerate(body):
        tag = child.tag.split('}')[1] if '}' in child.tag else child.tag
        if tag != 'p':
            continue
        texts = child.findall(f'{{{WPML}}}r/{{{WPML}}}t')
        text = ''.join(t.text or '' for t in texts).strip()
        if text == heading_text:
            return i
    return -1


def embed_all():
    doc = Document(str(DOCX_PATH))
    body = doc.element.body

    for heading_text, img_file, width_inches, caption in INSERTIONS:
        img_path = IMG_DIR / img_file
        if not img_path.exists():
            print(f'⚠️ 图片不存在: {img_path}')
            continue

        # Find target position
        idx = find_heading_index(body, heading_text)
        if idx < 0:
            print(f'⚠️ 未找到标题: {heading_text}')
            continue

        # 1) Add image to end of doc (registers blob + rId with doc)
        img_para = doc.add_paragraph()
        img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_para.paragraph_format.space_before = Pt(6)
        img_para.paragraph_format.space_after = Pt(0)
        run = img_para.add_run()
        run.add_picture(str(img_path), width=Inches(width_inches))

        # 2) Add caption to end of doc
        cap_para = doc.add_paragraph()
        cap_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_para.paragraph_format.space_before = Pt(2)
        cap_para.paragraph_format.space_after = Pt(10)
        r2 = cap_para.add_run(caption)
        r2.font.size = Pt(9)
        r2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        r2.font.italic = True

        # 3) Move elements: remove from end, insert after heading
        p_img_elem = img_para._element
        p_cap_elem = cap_para._element

        body.remove(p_img_elem)
        body.remove(p_cap_elem)
        body.insert(idx + 1, p_cap_elem)
        body.insert(idx + 1, p_img_elem)

        print(f'✅ {img_file}')

    doc.save(str(DOCX_PATH))
    print(f'\n✅ 已保存: {DOCX_PATH}')


if __name__ == '__main__':
    import subprocess
    print("Regenerating base document...")
    subprocess.run(['python', 'scripts/gen_docs.py'], cwd=str(Path(__file__).parent.parent), check=True)
    print()
    embed_all()
